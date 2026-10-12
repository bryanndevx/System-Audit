import os
from collections.abc import Iterator
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from fastapi import APIRouter, File, HTTPException, UploadFile, status
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel

SHEET_NAME = "Checklist Maestro"
EXPECTED_HEADERS = (
    "Fase",
    "#",
    "Paso / Actividad",
    "Entregable Esperado",
    "Prioridad",
    "Estado",
    "Responsable",
    "Notas / Evidencia",
)
DEFAULT_MAX_FILE_BYTES = 5 * 1024 * 1024
ALLOWED_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",
}

router = APIRouter(prefix="/api/checklists", tags=["checklists"])


class ActividadImportada(BaseModel):
    numero: str
    actividad: str
    entregable_esperado: str | None
    prioridad: str | None


class FaseImportada(BaseModel):
    nombre: str
    actividades: list[ActividadImportada]


class ChecklistImportado(BaseModel):
    nombre_archivo: str
    cantidad_fases: int
    cantidad_actividades: int
    fases: list[FaseImportada]


class ChecklistValidationError(ValueError):
    pass


def _max_file_bytes() -> int:
    raw_value = os.getenv("CHECKLIST_MAX_BYTES")
    if raw_value is None:
        return DEFAULT_MAX_FILE_BYTES

    try:
        value = int(raw_value)
    except ValueError:
        return DEFAULT_MAX_FILE_BYTES

    return value if value > 0 else DEFAULT_MAX_FILE_BYTES


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized(value: object) -> str:
    return " ".join(_text(value).split()).casefold()


def _find_data_rows(rows: Iterator[tuple[object, ...]]) -> tuple[int, Iterator[tuple[object, ...]]]:
    expected = tuple(header.casefold() for header in EXPECTED_HEADERS)

    for row_number, row in enumerate(rows, start=1):
        actual = tuple(_normalized(value) for value in row[: len(EXPECTED_HEADERS)])
        if actual == expected:
            return row_number, rows

    raise ChecklistValidationError(
        "No se encontró la fila de encabezados esperada en las columnas A:H. "
        "Verifique que incluya: " + ", ".join(EXPECTED_HEADERS) + "."
    )


def parse_checklist(content: bytes, file_name: str) -> ChecklistImportado:
    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (InvalidFileException, BadZipFile, OSError, ValueError) as exc:
        raise ChecklistValidationError(
            "El archivo no es un libro Excel .xlsx válido o está dañado."
        ) from exc

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ChecklistValidationError(
                f"El archivo debe contener una hoja llamada '{SHEET_NAME}'."
            )

        worksheet = workbook[SHEET_NAME]
        rows = worksheet.iter_rows(values_only=True)
        header_row, data_rows = _find_data_rows(rows)

        phases: list[FaseImportada] = []
        current_phase: FaseImportada | None = None
        current_phase_key = ""

        for row_number, row in enumerate(data_rows, start=header_row + 1):
            values = tuple(row[: len(EXPECTED_HEADERS)])
            padded_values = values + (None,) * (len(EXPECTED_HEADERS) - len(values))
            phase_value, number, activity, deliverable, priority, _, _, _ = padded_values

            if not any(_text(value) for value in padded_values):
                continue

            phase_name = _text(phase_value)
            if phase_name:
                phase_key = _normalized(phase_name)
                if current_phase is None or phase_key != current_phase_key:
                    current_phase = FaseImportada(nombre=phase_name, actividades=[])
                    phases.append(current_phase)
                    current_phase_key = phase_key

            activity_name = _text(activity)
            activity_number = _text(number)

            if not activity_name:
                if activity_number or _text(deliverable) or _text(priority):
                    raise ChecklistValidationError(
                        f"La fila {row_number} contiene datos de actividad, pero no tiene 'Paso / Actividad'."
                    )
                continue

            if current_phase is None:
                raise ChecklistValidationError(
                    f"La actividad de la fila {row_number} no tiene una fase asociada."
                )
            if not activity_number:
                raise ChecklistValidationError(
                    f"La actividad de la fila {row_number} no tiene valor en la columna '#'."
                )

            current_phase.actividades.append(
                ActividadImportada(
                    numero=activity_number,
                    actividad=activity_name,
                    entregable_esperado=_text(deliverable) or None,
                    prioridad=_text(priority) or None,
                )
            )

        phases_with_activities = [phase for phase in phases if phase.actividades]
        activity_count = sum(len(phase.actividades) for phase in phases_with_activities)
        if activity_count == 0:
            raise ChecklistValidationError(
                "La hoja 'Checklist Maestro' no contiene actividades válidas."
            )

        return ChecklistImportado(
            nombre_archivo=Path(file_name).name,
            cantidad_fases=len(phases_with_activities),
            cantidad_actividades=activity_count,
            fases=phases_with_activities,
        )
    finally:
        workbook.close()


@router.post(
    "/import",
    response_model=ChecklistImportado,
    status_code=status.HTTP_200_OK,
)
async def import_checklist(file: UploadFile = File(...)) -> ChecklistImportado:
    file_name = Path(file.filename or "").name
    if not file_name or Path(file_name).suffix.casefold() != ".xlsx":
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Seleccione un archivo con extensión .xlsx.",
        )

    if file.content_type and file.content_type not in ALLOWED_CONTENT_TYPES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El tipo de archivo no corresponde a un libro Excel .xlsx.",
        )

    max_bytes = _max_file_bytes()
    content = await file.read(max_bytes + 1)
    await file.close()

    if not content:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo está vacío.",
        )
    if len(content) > max_bytes:
        max_megabytes = max_bytes / (1024 * 1024)
        raise HTTPException(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            detail=f"El archivo supera el tamaño máximo permitido de {max_megabytes:g} MiB.",
        )

    try:
        return parse_checklist(content, file_name)
    except ChecklistValidationError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail=str(exc),
        ) from exc
